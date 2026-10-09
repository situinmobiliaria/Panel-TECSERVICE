# -*- coding: utf-8 -*-
"""Listado consolidado de clientes TECSERVICE para Marketing (una fila por cliente).
Lee data/CONTRATOS -FACTURACION -SATISFACCION -VISITAS.xlsb sin modificarlo y genera
Clientes_TECSERVICE_Marketing.xlsx en la raíz del proyecto.
Uso: python src/gen_clientes_marketing.py"""
import os, re, unicodedata, difflib
from collections import Counter, defaultdict
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.utils import get_column_letter as L

DIR = os.path.dirname(os.path.abspath(__file__)); ROOT = os.path.dirname(DIR)
XLSB = os.path.join(ROOT, "data", "CONTRATOS -FACTURACION -SATISFACCION -VISITAS.xlsb")
OUT = os.path.join(ROOT, "Clientes_TECSERVICE_Marketing.xlsx")
ANO = 2026


def sin_tilde(s):
    return "".join(c for c in unicodedata.normalize("NFD", str(s)) if unicodedata.category(c) != "Mn")


_LEGAL = {"SA", "S", "A", "P", "SPA", "LTDA", "LIMITADA", "EIRL", "CIA", "SAC", "AG"}


def clave(s):
    """Mayúsculas, sin tildes, sin puntuación ni espacios de más, sin sufijo societario."""
    t = re.sub(r"[^A-Z0-9]+", " ", sin_tilde(s).upper()).strip()
    tk = t.split()
    while len(tk) > 1 and tk[-1] in _LEGAL:
        tk.pop()
    return " ".join(tk)


# Mismas equivalencias revisadas a mano que usa extractor.py (_BI_EQUIV)
EQUIV = [("CESFAM JOSE BAUZA FRAU", "CESFAM DR. JOSE BAUZA FRAU"),
         ("CLINICA UNIVERSIDAD ANDES", "CLINICA UNIVERSIDAD LOS ANDES"),
         ("HOSPITAL DE LA FLORIDA", "HOSPITAL CLINICO METROPOLITANO LA FLORIDA"),
         ("INSTITUTO DE SEGURIDAD DEL TRABAJO", "HOSPITAL CLINICO IST VINA DEL MAR"),
         ("HOSPITAL LA LIGUA", "HOSPITAL SAN AGUSTIN"), ("HOSPITAL DE LIMACHE", "HOSPITAL SANTO TOMAS DE LIMACHE"),
         ("MEGASALUD CONCHALI", "REDSALUD CONCHALI"),
         # Errores de tipeo y variantes de escritura del mismo cliente (revisado 06-10-2026)
         ("COMERCIALIZADORA DE INSUMOS MEDICOS CLP SA", "COMERCIALIZADORA DEI NSUMOS MEDICOS CLP SA"),
         ("CESFAM ENTRE LAGOS", "CESFAM ENTRELAGOS"),
         ("I. MUNICIPALIDAD DE HUECHURABA", "MUNICIPALIDAD DE HUECHURABA"),
         ("CENTRO CLINICO MILITAR DE IQUIQUE", "CENTRO CLINICO MILITAR IQUIQUE"),
         ("REDMEDICAL LTDA", "RED MEDICAL"),
         ("SERVICIO DE SALUD ARAUCANIA NORTE", "SERVICIO SALUD ARAUCANIA NORTE"),
         ("CESFAM DR. PEDRO PULGAR MELGAREJO", "CESFAM PEDRO PULGAR MELGAREJO"),
         ("I. MUNICIPALIDAD DE MARIA ELENA", "MUNCIPALIDAD DE MARIA ELENA"),
         ("I. MUNICIPALIDAD DE RETIRO", "MUNCIPALIDAD DE RETIRO"),
         ("HOSPITAL DE ALTO HOSPICIO", "HOSPITAL ALTO HOSPICIO"),
         ("I.MUNICIPALIDAD DE CURICO DPTO. SALUD", "MUNICIPALIDAD DE CURICO DPTO. DE SALUD"),
         ("HOSPITAL BASE DE OSORNO", "HOSPITAL BASE OSORNO"),
         ("HOSPITAL SANTA JUANA", "HOSPITAL DE SANTA JUANA"),
         ("HOSPITAL SAN CARLOS", "HOSPITAL DE SAN CARLOS"),
         ("HOSPITAL SANTA CRUZ", "HOSPITAL DE SANTA CRUZ"),
         ("HOSPITAL LA SERENA", "HOSPITAL DE LA SERENA"),
         ("HOSPITAL LOS VILOS", "HOSPITAL DE LOS VILOS"),
         ("ESPINDOLA Y ASOCIADOS LTDA.", "ESPINOLA ASOCIADOS LTDA."),
         ("CENTRO OFTALMOLOGICO DR EDGARDO CARRENO SPA", "CENTRO OFTALMOLOGICO DR. E. CARRENO SPA"),
         ("CORPORACION MUNICIPAL DESARROLLO SOCIAL ANTOFAGASTA", "CORP. MUNICIPAL DESARROLLO SOCIAL ANTOFAGASTA"),
         ("CORPORACION MUNICIPAL DESARROLLO SOCIAL ANTOFAGASTA", "CORPORACION MUNICIPAL DESARROLLO SOCIAL ANTOFAGAST"),
         ("HOSPITAL DE LIMACHE", "HOSPITAL LIMACHE")]
CANON = {clave(b): clave(a) for a, b in EQUIV}


def K(s):
    k = clave(s)
    return CANON.get(k, k)


REG_ALIAS = {"LA ARAUCANIA": "Araucanía", "ARAUCANIA": "Araucanía", "BIOBIO": "Biobío", "BIO BIO": "Biobío",
             "METROPOLITANA": "Metropolitana", "METROPOLITANA DE SANTIAGO": "Metropolitana", "OHIGGINS": "O'Higgins",
             "MAGALLANES": "Magallanes", "MAGALLANES Y LA ANTARTICA CHILENA": "Magallanes", "NUBLE": "Ñuble",
             "VALPARAISO": "Valparaíso", "VALPARAISO VALPARAISO": "Valparaíso", "LOS LAGOS": "Los Lagos",
             "LOS RIOS": "Los Ríos", "AYSEN": "Aysén", "TARAPACA": "Tarapacá", "COQUIMBO": "Coquimbo",
             "MAULE": "Maule", "ATACAMA": "Atacama", "ANTOFAGASTA": "Antofagasta",
             "ARICA Y PARINACOTA": "Arica y Parinacota"}


def region(v):
    if v is None or (isinstance(v, float) and pd.isna(v)):
        return None
    k = re.sub(r"[^A-Z0-9 ]", "", sin_tilde(v).upper().replace("'", "")).strip()
    k = re.sub(r"\s+", " ", k)
    k = re.sub(r"^REGION( DE| DEL)? ", "", k)
    return REG_ALIAS.get(k)   # "Sin Información", 0, 0x2a… -> None


def num(s):
    return pd.to_numeric(s, errors="coerce").fillna(0.0)


print("Leyendo hojas…")
bi = pd.read_excel(XLSB, sheet_name="BASE INSTALADA", engine="pyxlsb")
fa = pd.read_excel(XLSB, sheet_name="FACTURACION", engine="pyxlsb")
lg = pd.read_excel(XLSB, sheet_name="Facturacion a Fecha", engine="pyxlsb")

# ── BASE INSTALADA: equipos activos (Habilitado=Si y no Fuera de servicio), igual que el panel ──
hab = bi.iloc[:, 0].astype(str).str.strip().str.upper() == "SI"
fuera = bi.iloc[:, 1].astype(str).str.strip().str.upper() == "SI"
b = bi[hab & ~fuera].copy()
n_activos = len(b)


def nombre_bi(r):
    for i in (26, 13):   # «Nombre Analisis» y, si está vacío, «Clasificación 1»
        v = r.iloc[i]
        if isinstance(v, str) and v.strip().upper() not in {"", "NONE", "VACIO", "VACÍO", "0", "0.0"}:
            return v.strip()
    return None


b["nom"] = b.apply(nombre_bi, axis=1)
n_vacio = int(b.nom.isna().sum())
b = b[b.nom.notna()]
es_gemco = b.nom.str.upper().str.contains("GEMCO")
n_gemco = int(es_gemco.sum())
b = b[~es_gemco].copy()
b["k"] = b.nom.map(K)
b["linea"] = b.iloc[:, 25].fillna("OTRO").astype(str).str.strip().str.upper().replace({"": "OTRO"})
b["pot"] = b.iloc[:, 27].astype(str).str.strip().str.upper().isin({"SI", "SÍ"})
b["reg"] = b.iloc[:, 29].map(region)

LKEY = {sin_tilde(l): l for l in set(b.linea)}
ORDEN = ["DENTAL", "ESTERILIZACION", "ENDOSCOPIA", "INCARDIA", "MOBILIARIO CLINICO", "MMQ", "REAS", "OTRO"]
LINEAS = [LKEY[x] for x in ORDEN if x in LKEY] + [l for l in sorted(set(b.linea)) if sin_tilde(l) not in ORDEN]
ETQ = {"DENTAL": "Dental", "ESTERILIZACION": "Esterilización", "ENDOSCOPIA": "Endoscopía", "INCARDIA": "Incardia",
       "MOBILIARIO CLINICO": "Mobiliario Clínico", "MMQ": "MMQ", "REAS": "REAS", "OTRO": "Otro"}
etq = lambda l: ETQ.get(sin_tilde(l), l.title())

clientes = {}


def cli(k):
    return clientes.setdefault(k, {"nombres": Counter(), "regs": Counter(), "bi": 0, "bi_pot": 0, "bl": Counter(),
                                    "bp": Counter(), "tot": 0.0, "con": 0.0, "fuente": set()})


for k, g in b.groupby("k"):
    d = cli(k); d["fuente"].add("BI")
    d["nombres"].update(g.nom); d["regs"].update(g.reg.dropna())
    d["bi"] = len(g); d["bi_pot"] = int(g.pot.sum())
    for l, gl in g.groupby("linea"):
        d["bl"][l] = len(gl); d["bp"][l] = int(gl.pot.sum())

# ── FACTURACION (hoja del panel): total 2026 y contratos 2026; otros = total − contratos ──
fa["tot"] = num(fa.iloc[:, 3]); fa["con"] = num(fa.iloc[:, 7])
fa["nom"] = fa.apply(lambda r: next((r.iloc[i] for i in (2, 0, 1) if isinstance(r.iloc[i], str) and r.iloc[i].strip()), None), axis=1)
fa = fa[fa.nom.notna()].copy()
fa["k"] = fa.nom.map(K)
alias_sheet = defaultdict(set)   # clave de cualquier nombre de la hoja -> clave de consolidación
for _, r in fa.iterrows():
    for i in (0, 1, 2):
        if isinstance(r.iloc[i], str):
            alias_sheet[K(r.iloc[i])].add(r["k"])
for k, g in fa.groupby("k"):
    d = cli(k); d["fuente"].add("FACTURACION")
    d["nombres"].update(g.nom); d["tot"] += g.tot.sum(); d["con"] += g.con.sum()
    d["regs"].update(g.iloc[:, 30].map(region).dropna())

# ── Facturas (Facturacion a Fecha): clientes con facturación 2026 que no están en la hoja FACTURACION ──
lc = lg.columns
t = lg[(lg[lc[43]].astype(str).str.strip() == "TS") & (num(lg[lc[40]]) == ANO)].copy()
t["monto"] = num(t[lc[18]])
sin_nombre = float(t[t[lc[7]].isna()].monto.sum())   # facturas sin nombre de cliente: no se pueden asignar
t = t[t[lc[7]].notna()].copy()
t["nom"] = t[lc[7]].astype(str).str.strip(); t["k0"] = t.nom.map(K)
ledger_total = float(t.monto.sum()) + sin_nombre
t["k"] = t.k0
solo = t[~t.k0.isin(alias_sheet)]
for k, g in solo.groupby("k"):
    d = cli(k); d["fuente"].add("FACTURAS")
    d["nombres"].update(g.nom); d["tot"] += g.monto.sum()    # sin contrato -> todo es "otros"
    d["regs"].update(g[lc[53]].map(region).dropna())
t_mapeado = float(t[t.k0.isin(alias_sheet)].monto.sum())

# ── Región de respaldo: hojas «Región Clientes» y «Región BI» (sólo cuando BI/FACTURACION no la traen) ──
reg_resp = {}
for hoja, cn, cr in (("Región Clientes", 0, 1), ("Región BI", 1, 4)):
    d_ = pd.read_excel(XLSB, sheet_name=hoja, engine="pyxlsb")
    for nom_, rg_ in zip(d_.iloc[:, cn], d_.iloc[:, cr]):
        if isinstance(nom_, str) and region(rg_):
            reg_resp.setdefault(K(nom_), region(rg_))
n_resp = 0

# ── Filas ──
filas = []
for k, d in clientes.items():
    nombre = d["nombres"].most_common(1)[0][0] if d["nombres"] else k
    reg = d["regs"].most_common(1)[0][0] if d["regs"] else reg_resp.get(k)
    if reg and not d["regs"]:
        n_resp += 1
    reg = reg or "Sin región"
    filas.append({"Cliente": nombre.strip().upper(), "Región": reg, "BI total": d["bi"], "BI Potencial ST": d["bi_pot"],
                  "Fact total": round(d["tot"], 2), "Fact contratos": round(d["con"], 2),
                  "Fact otros": round(d["tot"] - d["con"], 2), "_bl": d["bl"], "_bp": d["bp"],
                  "_f": "+".join(sorted(d["fuente"])), "_k": k})
filas.sort(key=lambda r: (sin_tilde(r["Cliente"]), r["_k"]))
# si dos claves distintas dieran el mismo nombre visible, se desambigua con la región
vis = Counter(sin_tilde(r["Cliente"]) for r in filas)
for r in filas:
    if vis[sin_tilde(r["Cliente"])] > 1:
        r["Cliente"] += f" ({r['Región']})"

# ── Excel ──
wb = Workbook(); ws = wb.active; ws.title = "Clientes TECSERVICE"
cab = ["Cliente", "Región", "Base instalada total", "Base instalada con mantenimiento potencial (Potencial ST = Sí)",
       "Facturación 2026 total", "Facturación 2026 contratos", "Facturación 2026 otros"]
for l in LINEAS:
    cab += [f"BI {etq(l)} total", f"BI {etq(l)} Potencial ST"]
ws.append(cab)
for r in filas:
    row = [r["Cliente"], r["Región"], r["BI total"], r["BI Potencial ST"], r["Fact total"], r["Fact contratos"], r["Fact otros"]]
    for l in LINEAS:
        row += [r["_bl"].get(l, 0), r["_bp"].get(l, 0)]
    ws.append(row)
n = len(filas) + 1; nc = len(cab)
CLP = '"$"#,##0;[Red]-"$"#,##0'
for c in range(3, nc + 1):
    for rr in range(2, n + 1):
        ws.cell(rr, c).number_format = CLP if 5 <= c <= 7 else "#,##0"
tab = Table(displayName="ClientesTECSERVICE", ref=f"A1:{L(nc)}{n}")
tab.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
ws.add_table(tab)
ws.freeze_panes = "B2"
ws.row_dimensions[1].height = 62
for c in range(1, nc + 1):
    ws.cell(1, c).alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
    ws.column_dimensions[L(c)].width = 46 if c == 1 else 18 if c == 2 else 17 if 5 <= c <= 7 else 14

# ── Validación ──
v = wb.create_sheet("Validación")
rng = lambda c: f"'Clientes TECSERVICE'!{L(c)}2:{L(c)}{n}"
src_bi, src_pot = int(len(b)), int(b.pot.sum())
fa_tot, fa_con = float(fa.tot.sum()), float(fa.con.sum())
solo_tot = float(solo.monto.sum())
rows = [("Control", "Planilla", "Fuente / esperado", "Diferencia", "Resultado")]


def chk(lbl, f, exp):
    r = len(rows) + 1
    rows.append((lbl, f, exp, f"=B{r}-C{r}", f'=IF(ABS(D{r})<0.5,"OK","REVISAR")'))


chk("Clientes únicos (filas)", f"=COUNTA({rng(1)})", len(filas))
chk("Facturación 2026 total", f"=SUM({rng(5)})", round(fa_tot + solo_tot, 2))
chk("Facturación 2026 contratos", f"=SUM({rng(6)})", round(fa_con, 2))
chk("Facturación 2026 otros", f"=SUM({rng(7)})", round(fa_tot - fa_con + solo_tot, 2))
chk("Total = contratos + otros", f"=SUM({rng(5)})", f"=SUM({rng(6)})+SUM({rng(7)})")
chk("Base instalada total (equipos)", f"=SUM({rng(3)})", src_bi)
chk("Base instalada Potencial ST = Sí", f"=SUM({rng(4)})", src_pot)
chk("BI total = suma de las líneas", f"=SUM({rng(3)})", "=" + "+".join(f"SUM({rng(8 + 2 * i)})" for i in range(len(LINEAS))))
chk("BI Potencial ST = suma de las líneas", f"=SUM({rng(4)})", "=" + "+".join(f"SUM({rng(9 + 2 * i)})" for i in range(len(LINEAS))))
for r_ in rows:
    v.append(list(r_))
for r in range(2, len(rows) + 1):
    for c in (2, 3, 4):
        v.cell(r, c).number_format = '"$"#,##0.00' if r in (3, 4, 5, 6) else "#,##0"
for c in range(1, 6):
    v.cell(1, c).font = Font(bold=True, color="FFFFFF"); v.cell(1, c).fill = PatternFill("solid", fgColor="1F4E78")
v.append([])
v.append(["Cómo se construyó la columna «Fuente / esperado»", "Valor"]); v.cell(v.max_row, 1).font = Font(bold=True)
info = [
    ("Equipos activos en BASE INSTALADA (Habilitado = Sí y no Fuera de servicio)", n_activos, "#,##0"),
    ("  menos equipos sin cliente identificable (VACIO)", n_vacio, "#,##0"),
    ("  menos equipos de cuentas internas GEMCO", n_gemco, "#,##0"),
    ("  = equipos en la planilla", src_bi, "#,##0"),
    ("Facturación 2026 de la hoja FACTURACION (col. D, un total por cliente)", fa_tot, '"$"#,##0'),
    ("  de la cual contratos (col. H)", fa_con, '"$"#,##0'),
    ("Facturación 2026 de clientes que sólo están en Facturacion a Fecha (empresa TS)", solo_tot, '"$"#,##0'),
    ("Total 2026 en Facturacion a Fecha (empresa TS, todas las líneas)", ledger_total, '"$"#,##0'),
    ("  de eso, facturas sin nombre de cliente (no incluidas en la planilla)", sin_nombre, '"$"#,##0'),
    ("  parte de eso que corresponde a clientes de la hoja FACTURACION", t_mapeado, '"$"#,##0'),
    ("  lo que la hoja FACTURACION declara para esos mismos clientes", fa_tot, '"$"#,##0'),
    ("  diferencia (la hoja FACTURACION es menor; ver supuestos)", t_mapeado - fa_tot, '"$"#,##0'),
]
for a, x, f in info:
    v.append([a, x]); v.cell(v.max_row, 2).number_format = f
v.column_dimensions["A"].width = 80
for c in "BCDE":
    v.column_dimensions[c].width = 22

v.append([]); v.append(["Posibles duplicados NO fusionados (nombres muy parecidos; revisar a mano)"]); v.cell(v.max_row, 1).font = Font(bold=True)
v.append(["Cliente A", "Cliente B", "Similitud"])
ks = [(r["Cliente"], r["_k"]) for r in filas]; pares = []
for i in range(len(ks)):
    a = ks[i][1]
    for j in range(i + 1, len(ks)):
        sm = difflib.SequenceMatcher(None, a, ks[j][1])
        if sm.real_quick_ratio() >= .9 and sm.quick_ratio() >= .9 and sm.ratio() >= .9:
            pares.append((ks[i][0], ks[j][0], round(sm.ratio(), 3)))
for p in sorted(pares, key=lambda x: -x[2])[:80]:
    v.append(list(p))

wb.save(OUT)
neg = sum(1 for r in filas if r["Fact otros"] < -0.5)
print(f"OK {OUT}\nclientes={len(filas)} BI={src_bi} pot={src_pot} lineas={LINEAS}\n"
      f"fact total={fa_tot + solo_tot:,.0f} contr={fa_con:,.0f} otros={fa_tot - fa_con + solo_tot:,.0f}"
      f" | solo facturas={solo_tot:,.0f} ({solo.k.nunique()} clientes) | ledger TS 2026={ledger_total:,.0f}"
      f" mapeado={t_mapeado:,.0f} | otros negativos={neg} | pares parecidos={len(pares)}")
print(Counter(r["_f"] for r in filas)); print(Counter(r["Región"] for r in filas).most_common(20))
print("region de respaldo", n_resp)
print("activos", n_activos, "vacio", n_vacio, "gemco", n_gemco)
